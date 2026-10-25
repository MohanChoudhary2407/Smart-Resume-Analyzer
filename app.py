import io
import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from collections import Counter

from flask import Flask, render_template, request, redirect, url_for, flash, send_file, abort
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
import spacy
from spacy.matcher import PhraseMatcher

BASE_DIR = Path(__file__).resolve().parent
DATABASE = BASE_DIR / "resume_analyzer.db"
ALLOWED_EXTENSIONS = {"pdf", "docx"}
MAX_UPLOAD_MB = 8

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-development-secret")
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

# Role skills are intentionally editable starter lists, not a definitive hiring standard.
ROLE_SKILLS = {
    "Data Analyst": [
        "Python", "SQL", "Excel", "Power BI", "Tableau", "Statistics",
        "Data Cleaning", "Pandas", "NumPy", "Data Visualization", "Communication"
    ],
    "Web Developer": [
        "HTML", "CSS", "JavaScript", "React", "Node.js", "REST API",
        "Git", "Responsive Design", "SQL", "Testing", "Accessibility"
    ],
    "AI Engineer": [
        "Python", "Machine Learning", "Deep Learning", "NLP", "Pandas",
        "NumPy", "Scikit-learn", "TensorFlow", "PyTorch", "Git", "SQL"
    ],
    "Cloud Engineer": [
        "Linux", "Networking", "AWS", "Azure", "Google Cloud", "Docker",
        "Kubernetes", "CI/CD", "Git", "Python", "Security"
    ],
}

SKILL_ALIASES = {
    "js": "JavaScript",
    "javascript": "JavaScript",
    "html5": "HTML",
    "css3": "CSS",
    "node": "Node.js",
    "nodejs": "Node.js",
    "powerbi": "Power BI",
    "scikit learn": "Scikit-learn",
    "sklearn": "Scikit-learn",
    "natural language processing": "NLP",
    "machine-learning": "Machine Learning",
    "ml": "Machine Learning",
    "deep-learning": "Deep Learning",
    "restful api": "REST API",
    "rest apis": "REST API",
    "gcp": "Google Cloud",
    "amazon web services": "AWS",
    "microsoft azure": "Azure",
}

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS analyses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                original_filename TEXT NOT NULL,
                target_role TEXT NOT NULL,
                resume_score INTEGER NOT NULL,
                ats_score INTEGER NOT NULL,
                matched_skills TEXT NOT NULL,
                missing_skills TEXT NOT NULL,
                suggestions TEXT NOT NULL
            )
        """)

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_pdf_text(file_stream):
    reader = PdfReader(file_stream)
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()

def extract_docx_text(file_stream):
    document = Document(file_stream)
    chunks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks).strip()

def normalize_text(text):
    text = text.replace("\u2019", "'").replace("\u2013", "-").replace("\u2014", "-")
    return re.sub(r"\s+", " ", text).strip()

def has_skill(text, skill):
    """Match a skill as a phrase, avoiding many accidental substring matches."""
    pattern = re.escape(skill)
    if skill.lower() in {"c++", "c#", ".net", "node.js", "ci/cd"}:
        return re.search(pattern, text, re.IGNORECASE) is not None
    return re.search(r"(?<![A-Za-z0-9])" + pattern + r"(?![A-Za-z0-9])", text, re.IGNORECASE) is not None

'''Lightweight spaCy NLP pipeline: no large model download is required.
 PhraseMatcher uses spaCy tokenization and case-insensitive phrase matching to
 detect canonical skills and their aliases in resume/job-description text.'''

NLP = spacy.blank("en")
SKILL_MATCHER = PhraseMatcher(NLP.vocab, attr="LOWER")
ALL_SKILLS = list(dict.fromkeys(skill for role_skills in ROLE_SKILLS.values() for skill in role_skills))
PHRASE_TO_CANONICAL = {}
_skill_patterns = []
for _canonical in ALL_SKILLS:
    _phrases = [_canonical]
    _phrases.extend(alias for alias, target in SKILL_ALIASES.items() if target.lower() == _canonical.lower())
    for _phrase in _phrases:
        _key = _phrase.lower().strip()
        PHRASE_TO_CANONICAL[_key] = _canonical
        _skill_patterns.append(NLP.make_doc(_phrase))
SKILL_MATCHER.add("SKILL", _skill_patterns)


def extract_skills_nlp(text):
    """Extract known skills from text using spaCy tokenization + PhraseMatcher."""
    doc = NLP(normalize_text(text))
    matches = SKILL_MATCHER(doc)
    found = set()
    for _, start, end in matches:
        phrase = doc[start:end].text.lower().strip()
        canonical = PHRASE_TO_CANONICAL.get(phrase)
        if canonical:
            found.add(canonical)
    # Stable order makes results easier to read and test.
    return [skill for skill in ALL_SKILLS if skill in found]


def detect_skills(text, skills):
    extracted = set(extract_skills_nlp(text))
    return [skill for skill in skills if skill in extracted]

def calculate_resume_score(text):
    lower = text.lower()
    checks = [
        ("Contact information", bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)) or bool(re.search(r"\b(?:\+?\d[\d\s().-]{7,}\d)\b", text))),
        ("Education section", bool(re.search(r"\b(education|academic background|qualification)\b", lower))),
        ("Skills section", bool(re.search(r"\b(skills|technical skills|core competencies|technologies)\b", lower))),
        ("Projects section", bool(re.search(r"\b(projects|personal projects|academic projects)\b", lower))),
        ("Experience section", bool(re.search(r"\b(experience|internship|work history|employment)\b", lower))),
        ("Summary or objective", bool(re.search(r"\b(summary|objective|profile|about me)\b", lower))),
        ("Certifications section", bool(re.search(r"\b(certifications|certificates|licenses)\b", lower))),
        ("Action-oriented content", bool(re.search(r"\b(developed|built|created|implemented|designed|analyzed|improved|automated|led|delivered)\b", lower))),
        ("Measurable achievements", bool(re.search(r"\b\d+(\.\d+)?\s*(%|percent|users|customers|hours|days|projects|records|team members)\b", lower))),
        ("Readable length", 250 <= len(text.split()) <= 1200),
    ]
    earned = sum(weight for _, passed, weight in [(name, ok, 10) for name, ok in checks] if passed)
    breakdown = [{"name": name, "passed": passed, "points": 10 if passed else 0} for name, passed in checks]
    return min(100, earned), breakdown

def ats_format_checks(text):
    checks = []
    checks.append(("Contact details detected", bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text))))
    checks.append(("Common section headings", sum(bool(re.search(r"\b" + re.escape(h) + r"\b", text, re.I)) for h in ["education", "skills", "projects", "experience"]) >= 2))
    checks.append(("Text could be extracted", len(text.strip()) >= 150))
    checks.append(("No excessive decorative symbols", len(re.findall(r"[★◆●■►✓✔]", text)) <= 5))
    checks.append(("Reasonable document length", 150 <= len(text.split()) <= 1500))
    return checks

def build_suggestions(text, resume_score, missing_skills, format_checks, score_breakdown):
    suggestions = []
    failed_sections = [item["name"] for item in score_breakdown if not item["passed"]]
    if "Contact information" in failed_sections:
        suggestions.append("Add a clearly visible email address and/or phone number.")
    if "Education section" in failed_sections:
        suggestions.append("Add an Education section with degree, institution, and expected graduation year.")
    if "Skills section" in failed_sections:
        suggestions.append("Add a concise Skills section containing skills you can confidently demonstrate.")
    if "Projects section" in failed_sections:
        suggestions.append("Include 2–3 relevant projects, your contribution, tools used, and outcomes.")
    if "Experience section" in failed_sections:
        suggestions.append("If applicable, add internships, volunteering, freelance work, or relevant practical experience.")
    if "Summary or objective" in failed_sections:
        suggestions.append("Consider a short role-focused professional summary; avoid generic claims.")
    if "Certifications section" in failed_sections:
        suggestions.append("Add relevant certifications only if you have earned them.")
    if "Action-oriented content" in failed_sections:
        suggestions.append("Start bullet points with clear action verbs such as built, analyzed, automated, or implemented.")
    if "Measurable achievements" in failed_sections:
        suggestions.append("Where truthful, quantify impact (for example, records processed, time saved, or users supported).")
    if "Readable length" in failed_sections:
        suggestions.append("Review resume length: use concise, relevant bullet points and remove repeated content.")
    if missing_skills:
        preview = ", ".join(missing_skills[:6])
        suggestions.append("For the selected role, consider adding evidence of these skills if you genuinely have them: " + preview + ".")
    failed_format = [name for name, passed in format_checks if not passed]
    if failed_format:
        suggestions.append("Review ATS readability: " + "; ".join(failed_format).lower() + ".")
    if not suggestions:
        suggestions.append("Your resume passed the current basic checks. Tailor it to each job description and verify that every claim is accurate.")
    suggestions.append("Do not add a skill, certification, or achievement unless it accurately reflects your experience.")
    return suggestions

def analyze_resume(text, target_role, job_description=""):
    text = normalize_text(text)
    if len(text) < 80:
        raise ValueError("Not enough readable text was found. If this is a scanned PDF, use a text-based PDF or convert it to DOCX.")
    resume_score, score_breakdown = calculate_resume_score(text)
    role_required = ROLE_SKILLS[target_role]
    resume_skills = extract_skills_nlp(text)
    if job_description.strip():
        jd_skills = extract_skills_nlp(job_description)
        # Prefer the actual job description when it contains recognizable skills;
        # otherwise fall back to the selected role's starter list.
        required = jd_skills if jd_skills else role_required
        comparison_source = "Provided job description" if jd_skills else "Role skill list (job description had no recognized skills)"
    else:
        required = role_required
        comparison_source = "Role skill list"
    matched = [skill for skill in required if skill in resume_skills]
    missing = [skill for skill in required if skill not in resume_skills]
    ats_score = round((len(matched) / len(required)) * 100) if required else 0
    format_checks = ats_format_checks(text)
    suggestions = build_suggestions(text, resume_score, missing, format_checks, score_breakdown)
    return {
        "resume_score": resume_score,
        "ats_score": ats_score,
        "score_breakdown": score_breakdown,
        "format_checks": format_checks,
        "matched_skills": matched,
        "missing_skills": missing,
        "required_skills": required,
        "resume_skills": resume_skills,
        "job_description_skills": extract_skills_nlp(job_description) if job_description.strip() else [],
        "comparison_source": comparison_source,
        "suggestions": suggestions,
        "word_count": len(text.split()),
    }

@app.route("/", methods=["GET"])
def index():
    return render_template("index.html", roles=list(ROLE_SKILLS.keys()), max_upload_mb=MAX_UPLOAD_MB)

@app.route("/analyze", methods=["POST"])
def analyze():
    uploaded = request.files.get("resume")
    target_role = request.form.get("target_role", "")
    if not uploaded or not uploaded.filename:
        flash("Choose a PDF or DOCX resume to analyze.", "error")
        return redirect(url_for("index"))
    if not allowed_file(uploaded.filename):
        flash("Unsupported file type. Please upload a PDF or DOCX file.", "error")
        return redirect(url_for("index"))
    if target_role not in ROLE_SKILLS:
        flash("Choose a valid target role.", "error")
        return redirect(url_for("index"))

    filename = secure_filename(uploaded.filename) or "resume"
    try:
        uploaded.stream.seek(0)
        if filename.lower().endswith(".pdf"):
            text = extract_pdf_text(uploaded.stream)
        else:
            text = extract_docx_text(uploaded.stream)
        job_description = request.form.get("job_description", "").strip()
        if len(job_description) > 20000:
            flash("Job description is too long. Please keep it under 20,000 characters.", "error")
            return redirect(url_for("index"))
        result = analyze_resume(text, target_role, job_description)
    except Exception as exc:
        app.logger.exception("Resume analysis failed")
        flash("Could not read this resume. Check that the file is valid and not password-protected.", "error")
        return redirect(url_for("index"))

    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    # The raw resume text and uploaded file are not stored in SQLite.
    with get_db() as conn:
        cursor = conn.execute("""
            INSERT INTO analyses
            (created_at, original_filename, target_role, resume_score, ats_score,
             matched_skills, missing_skills, suggestions)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            created_at, filename, target_role, result["resume_score"], result["ats_score"],
            "||".join(result["matched_skills"]), "||".join(result["missing_skills"]),
            "||".join(result["suggestions"])
        ))
        analysis_id = cursor.lastrowid

    result.update({
        "analysis_id": analysis_id,
        "created_at": created_at,
        "filename": filename,
        "target_role": target_role,
    })
    return render_template("results.html", result=result)

@app.route("/history")
def history():
    with get_db() as conn:
        rows = conn.execute("SELECT * FROM analyses ORDER BY id DESC LIMIT 50").fetchall()
    items = []
    for row in rows:
        items.append({
            "id": row["id"],
            "created_at": row["created_at"],
            "filename": row["original_filename"],
            "target_role": row["target_role"],
            "resume_score": row["resume_score"],
            "ats_score": row["ats_score"],
        })
    return render_template("history.html", items=items)

@app.route("/report/<int:analysis_id>")
def download_report(analysis_id):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM analyses WHERE id = ?", (analysis_id,)).fetchone()
    if row is None:
        abort(404)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=0.65*inch, leftMargin=0.65*inch,
                            topMargin=0.65*inch, bottomMargin=0.65*inch)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Smart Resume Analyzer", styles["Title"]),
        Paragraph("Analysis Report", styles["Heading2"]),
        Spacer(1, 10),
        Paragraph(f"<b>File:</b> {row['original_filename']}", styles["BodyText"]),
        Paragraph(f"<b>Target role:</b> {row['target_role']}", styles["BodyText"]),
        Paragraph(f"<b>Analyzed:</b> {row['created_at']}", styles["BodyText"]),
        Spacer(1, 12),
        Paragraph(f"Resume completeness score: <b>{row['resume_score']}/100</b>", styles["Heading3"]),
        Paragraph(f"Role keyword match score: <b>{row['ats_score']}/100</b>", styles["Heading3"]),
        Spacer(1, 8),
        Paragraph("Matched role skills", styles["Heading3"]),
        Paragraph((", ".join(filter(None, row["matched_skills"].split("||"))) or "No listed skills matched") .replace("&", "&amp;"), styles["BodyText"]),
        Spacer(1, 8),
        Paragraph("Skills to review", styles["Heading3"]),
        Paragraph((", ".join(filter(None, row["missing_skills"].split("||"))) or "No missing skills from the current starter list") .replace("&", "&amp;"), styles["BodyText"]),
        Spacer(1, 8),
        Paragraph("Suggestions", styles["Heading3"]),
    ]
    for suggestion in filter(None, row["suggestions"].split("||")):
        safe = suggestion.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph("• " + safe, styles["BodyText"]))
        story.append(Spacer(1, 4))
    story += [
        Spacer(1, 12),
        Paragraph("Important: Scores are heuristic estimates for learning and resume improvement. They are not an actual ATS result or a hiring prediction. Skill lists are editable examples and may not reflect a specific job description.", styles["Italic"])
    ]
    doc.build(story)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=f"resume-analysis-{analysis_id}.pdf", mimetype="application/pdf")

@app.errorhandler(413)
def too_large(_error):
    flash(f"File too large. Maximum upload size is {MAX_UPLOAD_MB} MB.", "error")
    return redirect(url_for("index"))

init_db()

if __name__ == "__main__":
    # Debug mode is for local development only. Disable it for deployment.
    app.run(debug=True)
